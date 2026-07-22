from importlib.metadata import files
from pathlib import Path
from django import forms
from django.core.validators import RegexValidator

from django.shortcuts import render, HttpResponse, redirect
from django.views.decorators.http import require_GET,require_POST

from app01.utils.bootstrap import BootstrapForm
from app01 import models

class UploadForm(BootstrapForm):
    bootstrap_exclude_fields = ['avatar']
    nickname = forms.CharField(label='昵称', max_length=32)
    age = forms.IntegerField(label='年龄')
    gender_choices = (
        (1, '男'),
        (2, '女')
    )
    gender = forms.TypedChoiceField(
        label='性别',
        choices=gender_choices,
        coerce=int,
    )
    email = forms.CharField(label='邮箱',max_length=64,validators=[RegexValidator(r"(^[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+$)")])
    avatar = forms.FileField(label='头像')

@require_GET
def upload_lst(req):
    title = "个人信息"
    form = UploadForm()
    return render(req, 'upload_lst.html',{'form':form,'title':title})
    # if req.method == 'GET':
    #     return render(req, 'upload_lst.html',{'form':form,'title':title})
    # print(req.POST)  # 请求体中的内容
    # print((req.FILES))  # 请求发过来的文件
    # # 根据文件name属性获得文件对象
    # file_obj = req.FILES.get("avatar")

    with open(file_obj.name, mode='wb') as f:
        for chunk in file_obj.chunks():
            f.write(chunk)
    return HttpResponse('...')

@require_POST
def upload_form(req):
    form=UploadForm(data=req.POST,files=req.FILES)
    if form.is_valid():
        avatar_obj=form.cleaned_data.get('avatar')
        file_path=Path('media')/'avatar'/avatar_obj.name
        with open(file_path,mode='wb') as f:
            for chunk in avatar_obj.chunks():
                f.write(chunk)
        models.ProfileInfo.objects.create(
            nickname=form.cleaned_data['nickname'],
            age=form.cleaned_data['age'],
            gender=form.cleaned_data['gender'],
            email=form.cleaned_data['email'],
            avatar=file_path
        )
        return HttpResponse("上传成功")
    return render(req,'upload_lst',{'form':form,'title':'个人信息'})

def upload_multi(req):
    """ 批量上传（以批量添加到部门表为例） """
    # 获取excel文件对象
    file_obj = req.FILES.get('excel')
    print(type(file_obj))

    # 将对象传递给openpyxl，并有它读取文件内容
    from openpyxl import load_workbook
    wb = load_workbook(file_obj)
    sheet = wb.worksheets[0]  # 获取第一单元页

    # 获得excel页的第一行第一列对象，并打印对象值
    cell = sheet.cell(1, 1)
    print(cell.value)

    # 对于单元页中的每一行进行循环取值
    for row in sheet.iter_rows(min_row=2):
        # 每一行的第一个对象的值获取
        val = row[0].value
        print(val)
        # 判断部门表中是否存在该值（部门）
        exists = models.Department.objects.filter(title=val).exists()
        # 如果不存在则添加
        if not exists:
            models.Department.objects.create(title=val)
    return redirect('/depart/lst/')


