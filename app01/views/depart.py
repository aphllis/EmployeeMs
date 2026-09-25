import traceback

from django.shortcuts import render, redirect, HttpResponse
from django.template.defaultfilters import title
from django.views.decorators.http import require_POST
from django.http import JsonResponse
from app01 import models
from app01.utils.pagination import Pagination
from app01.utils.form import DepartModelForm
from openpyxl import load_workbook


# Create your views here.
def depart_lst(req):
    """ 部门列表 """
    # 数据库中获取所有部门信息
    # queryset是对象列表[对象，对象，对象]
    print(req.user)
    queryset = models.Department.objects.all()
    if not queryset:return render(req,"depart_lst.html")
    page_object = Pagination(request=req, queryset=queryset)
    context = {
        'queryset': page_object.page_queryset,
        'page_str': page_object.html()[0],
        'goto_page_str': page_object.html()[1],
        'start_index': page_object.start
    }
    return render(req, 'depart_lst.html', context)


def depart_add(req):
    """ 添加部门 """
    if req.method == 'GET':
        form = DepartModelForm()
        return render(req, 'change.html', {'form': form, 'title': '添加部门'})
    form = DepartModelForm(data=req.POST)
    if form.is_valid():
        form.save()
        return redirect('/depart/lst/')
    return render(req, 'change.html', {'form': form, 'title': '添加部门'})
    # 非ModalForm的数据保存
    # 获取表单提交的数据
    # title = req.POST.get('title')
    # # 保存到数据库
    # models.Department.objects.create(title=title)
    # # 重定向到部门列表
    # return redirect('/depart/lst/')


@require_POST
def depart_batch(req):
    file_obj = req.FILES.get('excel_file')
    print(type(file_obj))
    print('hello')
    if not file_obj:
        return JsonResponse({'error': '未收到文件'}, status=400)
    if not file_obj.name.endswith(('.xlsx', 'xls')):
        return JsonResponse({'error': '仅支持.xlsx或.xls格式'},status=400)
    try:
        wb=load_workbook(file_obj,data_only=True)
        sheet=wb.worksheets[0]
        rows=[list(row) for row in sheet.iter_rows(values_only=True)]
        for row in sheet.iter_rows(min_row=2):
            val=row[0].value
            exist=models.Department.objects.filter(title=val).exists()
            if not exist:
                models.Department.objects.create(title=val)
        return JsonResponse({
            'message':'上传成功',
            'rows_pre':rows[:5],
            'total_rows':len(rows)
        })
    except Exception as e:
        traceback.print_exc()
        return JsonResponse({'error':f'文件解析失败:{str(e)}'},status=400)

    return HttpResponse("上传")


def depart_del(req):
    """ 删除部门 """
    nid = req.GET.get('nid')
    models.Department.objects.filter(id=nid).delete()
    return redirect('/depart/lst/')


def depart_edit(req):
    """ 编辑部门 """
    nid = req.GET.get('nid')
    row_obj = models.Department.objects.filter(id=nid).first()
    if not row_obj:
        return redirect('/depart/lst/')
    if req.method == 'GET':
        form = DepartModelForm(instance=row_obj)
        return render(req, 'change.html', {'form': form, 'title': '部门编辑'})
    form = DepartModelForm(data=req.POST, instance=row_obj)
    if form.is_valid():
        form.save()
        return redirect('/depart/lst/')
    return render(req, 'change.html', {'form': form, 'title': '部门编辑'})
    # title = req.POST.get('title')
    # models.Department.objects.filter(id=nid).update(title=title)
    # return redirect('/depart/lst/')
